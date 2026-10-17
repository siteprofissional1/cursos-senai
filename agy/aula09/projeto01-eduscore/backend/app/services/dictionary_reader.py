"""
Serviço Leitor do Dicionário Oficial de Microdados do INEP
Lê a planilha Dicionário_Microdados_Enem_2025.xlsx para consulta de metadados
Padrão SENAI - Clean Code & Arquitetura de Dados
"""

from pathlib import Path
from typing import Dict, Any, List
import openpyxl
from backend.app.core.config import DATA_DIR

DICT_PATH = DATA_DIR / "microdados_enem_2025" / "DICIONÁRIO" / "Dicionário_Microdados_Enem_2025.xlsx"

class DictionaryReader:
    _cache: Dict[str, List[Dict[str, Any]]] = {}

    @classmethod
    def load_dictionary(cls) -> Dict[str, List[Dict[str, Any]]]:
        if cls._cache:
            return cls._cache

        if not DICT_PATH.exists():
            return {}

        wb = openpyxl.load_workbook(str(DICT_PATH), read_only=True, data_only=True)
        result = {}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            variables = []
            current_var = None

            for row in ws.iter_rows(values_only=True):
                col0 = row[0]
                col1 = row[1] if len(row) > 1 else None
                col2 = row[2] if len(row) > 2 else None
                col3 = row[3] if len(row) > 3 else None

                if col0 and str(col0).strip() and not str(col0).startswith("DICIONÁRIO") and not str(col0).startswith("NOME"):
                    var_name = str(col0).strip()
                    desc = str(col1).strip() if col1 else ""
                    current_var = {
                        "nome": var_name,
                        "descricao": desc,
                        "categorias": {}
                    }
                    variables.append(current_var)

                if current_var and col2 is not None and col3 is not None:
                    cat_key = str(col2).strip()
                    cat_val = str(col3).strip()
                    current_var["categorias"][cat_key] = cat_val

            result[sheet_name] = variables

        cls._cache = result
        return cls._cache
